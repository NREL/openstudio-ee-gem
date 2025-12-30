"""
Comprehensive workflow script for window retrofit analysis.

This script:
1. Loads a baseline model
2. Extracts baseline embodied carbon and runs baseline EnergyPlus simulation
3. Applies window_enhancement measure with EC3 embodied carbon calculations
4. Runs modified model EnergyPlus simulation
5. Calculates deltas (operational energy, embodied carbon, cost)
6. Updates optimization.xlsx with all scenario data
7. Generates spider chart for comparison

Usage:
    python run_comprehensive_workflow.py
"""

import sys
import os
from pathlib import Path
import json
import shutil
import subprocess

# Add measure directories to Python path
measure_dir = Path(__file__).parent.absolute()
window_measure_dir = measure_dir.parent / "window_enhancement"

# Import OpenStudio
try:
    import openstudio
except ImportError:
    print("Error: OpenStudio Python bindings not found.")
    print("Make sure OpenStudio is installed and Python bindings are available.")
    sys.exit(1)

# Import required modules
import pandas as pd
from openpyxl import load_workbook
import plotly.graph_objects as go
from dotenv import load_dotenv

# Import ECReport from the current measure directory
sys.path.insert(0, str(measure_dir))
from measure import ECReport
from call_RSmeans import RSMeansAPIClient

# Import WindowEnhancement from the window_enhancement measure directory
sys.path.insert(0, str(window_measure_dir))
import importlib.util
spec = importlib.util.spec_from_file_location("window_measure", window_measure_dir / "measure.py")
window_measure_module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(window_measure_module)
WindowEnhancement = window_measure_module.WindowEnhancement


def find_openstudio_cli():
    """Find OpenStudio CLI executable."""
    cli_paths = [
        r"C:\openstudio-3.8.0\bin\openstudio.exe",
        r"C:\openstudio-3.7.0\bin\openstudio.exe",
        r"C:\openstudio\bin\openstudio.exe",
        "openstudio",  # Try system PATH
    ]
    
    for path in cli_paths:
        try:
            result = subprocess.run([path, "--version"], 
                                  capture_output=True, text=True, timeout=5)
            if result.returncode == 0:
                print(f"✓ Found OpenStudio CLI: {path}")
                return path
        except (FileNotFoundError, subprocess.TimeoutExpired):
            continue
    
    return None


def run_energyplus_simulation(osm_path, weather_file, output_dir):
    """
    Run EnergyPlus simulation and return parsed results.
    
    Args:
        osm_path: Path to OSM model file
        weather_file: Path to EPW weather file
        output_dir: Directory for simulation output
    
    Returns:
        dict: Simulation results or None if failed
    """
    print(f"\n{'='*80}")
    print(f"Running EnergyPlus simulation: {osm_path.name}")
    print(f"{'='*80}\n")
    
    output_dir.mkdir(parents=True, exist_ok=True)
    
    # Create workflow JSON
    workflow_dict = {
        "seed_file": osm_path.name,
        "weather_file": weather_file.name,
        "measure_paths": [],
        "file_paths": [],
        "run_directory": "./"
    }
    
    osw_path = output_dir / "workflow.osw"
    with open(osw_path, 'w') as f:
        json.dump(workflow_dict, f, indent=2)
    
    # Copy files
    shutil.copy(osm_path, output_dir / osm_path.name)
    if weather_file.exists():
        shutil.copy(weather_file, output_dir / weather_file.name)
    
    # Find OpenStudio CLI
    cli_exe = find_openstudio_cli()
    if not cli_exe:
        print("✗ Could not find OpenStudio CLI. Skipping simulation.")
        return None
    
    # Run simulation
    print(f"Running simulation (this may take a few minutes)...")
    try:
        result = subprocess.run(
            [cli_exe, "run", "-w", str(osw_path)],
            cwd=str(output_dir),
            capture_output=True,
            text=True,
            timeout=300
        )
        
        if result.returncode == 0:
            print("✓ Simulation completed successfully")
        else:
            print(f"✗ Simulation failed with return code {result.returncode}")
            return None
    except subprocess.TimeoutExpired:
        print("✗ Simulation timed out")
        return None
    except Exception as e:
        print(f"✗ Error running simulation: {e}")
        return None
    
    # Parse results from eplustbl.html
    html_paths = [
        output_dir / "reports" / "eplustbl.html",
        output_dir / "run" / "eplustbl.html",
    ]
    
    for html_path in html_paths:
        if html_path.exists():
            print(f"✓ Found EnergyPlus report: {html_path}")
            return parse_energyplus_html(html_path)
    
    print("✗ eplustbl.html not found")
    return None


def parse_energyplus_html(html_path):
    """Parse EnergyPlus HTML report for energy metrics."""
    import re
    
    html_text = html_path.read_text(encoding="utf-8", errors="ignore")
    
    def extract_field(pattern):
        m = re.search(pattern, html_text, flags=re.IGNORECASE | re.DOTALL)
        return float(m.group(1).strip()) if m else 0.0
    
    data = {
        "total_site_energy_GJ": extract_field(r'Total Site Energy</td>\s*<td[^>]*>\s*([0-9.]+)'),
        "net_site_energy_GJ": extract_field(r'Net Site Energy</td>\s*<td[^>]*>\s*([0-9.]+)'),
        "total_source_energy_GJ": extract_field(r'Total Source Energy</td>\s*<td[^>]*>\s*([0-9.]+)'),
        "total_building_area_m2": extract_field(r'Total Building Area</td>\s*<td[^>]*>\s*([0-9.]+)'),
    }
    
    return data


def extract_embodied_carbon(model):
    """Extract total embodied carbon from model AdditionalProperties."""
    total_carbon = 0.0
    carbon_by_component = {}
    
    for construction in model.getConstructions():
        props = construction.additionalProperties()
        
        # Look for embodied carbon properties
        if props.hasFeature("total_embodied_carbon_kgCO2eq"):
            value = props.getFeatureAsDouble("total_embodied_carbon_kgCO2eq")
            if value.is_initialized():
                carbon = value.get()
                total_carbon += carbon
                carbon_by_component[construction.nameString()] = carbon
    
    return total_carbon, carbon_by_component


def apply_window_enhancement_measure(model, arguments):
    """
    Apply window enhancement measure to the model.
    
    Args:
        model: OpenStudio model
        arguments: dict of measure arguments
    
    Returns:
        bool: Success status
    """
    print(f"\n{'='*80}")
    print("Applying Window Enhancement Measure")
    print(f"{'='*80}\n")
    
    # Create OSRunner
    osw = openstudio.WorkflowJSON()
    runner = openstudio.measure.OSRunner(osw)
    
    # Create measure
    measure = WindowEnhancement()
    
    # Get measure arguments
    args = measure.arguments(model)
    arg_map = openstudio.measure.convertOSArgumentVectorToMap(args)
    
    # Set argument values
    for arg in args:
        arg_name = arg.name()
        if arg_name in arguments:
            value = arguments[arg_name]
            
            # Set value based on argument type
            if isinstance(value, bool):
                arg.setValue(value)
            elif isinstance(value, int):
                arg.setValue(int(value))
            elif isinstance(value, float):
                arg.setValue(float(value))
            else:
                arg.setValue(str(value))
            
            # Update the map
            arg_map[arg_name] = arg
            print(f"  Set {arg_name} = {value}")
    
    # Run measure
    result = measure.run(model, runner, arg_map)
    
    # Print results
    print(f"Result: {runner.result().value().valueName()}")
    
    for info in runner.result().info():
        print(f"INFO: {info.logMessage()}")
    
    for warning in runner.result().warnings():
        print(f"WARNING: {warning.logMessage()}")
    
    for error in runner.result().errors():
        print(f"ERROR: {error.logMessage()}")
    
    return result


def get_rsmeans_cost(component_type="windows", runner=None):
    """
    Get cost estimate from RSMeans API.
    
    Args:
        component_type: Type of component (e.g., 'windows')
        runner: Optional OSRunner for logging
    
    Returns:
        float: Cost estimate or 0.0 if unavailable
    """
    load_dotenv()
    client_id = os.getenv('client_id')
    client_secret = os.getenv('client_secret')
    
    if not client_id or not client_secret:
        print("⚠ RSMeans API credentials not found in .env file")
        return 0.0
    
    client = RSMeansAPIClient(client_id, client_secret, use_sandbox=True)
    
    if not client.authenticate():
        print("⚠ Failed to authenticate with RSMeans API")
        return 0.0
    
    print("✓ Successfully authenticated with RSMeans API")
    
    # Search for relevant cost line items
    search_results = client.search_unit_costlines(
        release_id='2019-an',
        measurement_system='imp',
        searchTerm=component_type
    )
    
    if search_results and len(search_results) > 0:
        # Get first result's unit cost (simplified - should be refined)
        try:
            cost = float(search_results[0].get('unitCost', 0.0))
            print(f"✓ Retrieved cost estimate: ${cost:.2f}")
            return cost
        except (ValueError, KeyError, IndexError):
            print("⚠ Could not parse cost from RSMeans results")
            return 0.0
    
    return 0.0


def update_optimization_spreadsheet(baseline_data, modified_data, delta_data, output_path):
    """
    Update optimization.xlsx with scenario data.
    
    Args:
        baseline_data: dict with baseline metrics
        modified_data: dict with modified metrics  
        delta_data: dict with delta metrics
        output_path: Path to save updated spreadsheet
    """
    print(f"\n{'='*80}")
    print("Updating Optimization Spreadsheet")
    print(f"{'='*80}\n")
    
    excel_path = measure_dir / 'resources' / 'optimization.xlsx'
    wb = load_workbook(excel_path, data_only=False, keep_links=True)
    
    if 'values' not in wb.sheetnames:
        print("✗ Sheet 'values' not found in Excel file")
        return
    
    ws = wb['values']
    
    # Map factor names to values
    factor_mapping = {
        'Embodied Carbon': {
            'Baseline': baseline_data['embodied_carbon'],
            'Scenario_1': modified_data['embodied_carbon'],
            'Delta': delta_data['embodied_carbon']
        },
        'Operational Energy': {
            'Baseline': baseline_data['operational_energy'],
            'Scenario_1': modified_data['operational_energy'],
            'Delta': delta_data['operational_energy']
        },
        'Cost': {
            'Baseline': baseline_data['cost'],
            'Scenario_1': modified_data['cost'],
            'Delta': delta_data['cost']
        }
    }
    
    # Update cells
    for row in ws.iter_rows(min_row=2):  # Skip header
        factor_name = str(row[0].value).strip() if row[0].value else ""
        
        if factor_name in factor_mapping:
            # Find column headers
            for col_idx, cell in enumerate(ws[1], start=1):
                col_header = str(cell.value).strip() if cell.value else ""
                
                if col_header in factor_mapping[factor_name]:
                    value = factor_mapping[factor_name][col_header]
                    ws.cell(row=row[0].row, column=col_idx, value=value)
                    print(f"  Updated {factor_name} - {col_header}: {value:.2f}")
    
    wb.save(output_path)
    print(f"\n✓ Spreadsheet saved: {output_path}")


def generate_spider_chart(baseline_data, modified_data, output_path):
    """Generate spider chart comparing baseline and modified scenarios."""
    print(f"\n{'='*80}")
    print("Generating Spider Chart")
    print(f"{'='*80}\n")
    
    categories = ['Embodied Carbon\n(kg CO2eq)', 
                  'Operational Energy\n(GJ)', 
                  'Cost\n($)']
    
    # Normalize values (as % of baseline)
    baseline_values = [
        baseline_data['embodied_carbon'],
        baseline_data['operational_energy'],
        baseline_data['cost']
    ]
    
    modified_values = [
        modified_data['embodied_carbon'],
        modified_data['operational_energy'],
        modified_data['cost']
    ]
    
    normalized_baseline = [100, 100, 100]  # Baseline is 100%
    normalized_modified = [
        (mod / base * 100) if base > 0 else 100
        for mod, base in zip(modified_values, baseline_values)
    ]
    
    fig = go.Figure()
    
    fig.add_trace(go.Scatterpolar(
        r=normalized_baseline,
        theta=categories,
        fill='toself',
        name='Baseline',
        line_color='blue'
    ))
    
    fig.add_trace(go.Scatterpolar(
        r=normalized_modified,
        theta=categories,
        fill='toself',
        name='Window Enhancement',
        line_color='red'
    ))
    
    fig.update_layout(
        polar=dict(
            radialaxis=dict(
                visible=True,
                range=[0, 150]
            )),
        showlegend=True,
        title="Window Retrofit Impact Comparison<br>(Normalized to Baseline = 100%)"
    )
    
    fig.write_html(output_path)
    print(f"✓ Spider chart saved: {output_path}")
    
    # Also show in browser
    fig.show()


def main():
    print(f"\n{'='*80}")
    print("COMPREHENSIVE WINDOW RETROFIT WORKFLOW")
    print(f"{'='*80}\n")
    
    # Configuration
    baseline_model_path = measure_dir / "tests" / "example_model.osm"
    weather_file = measure_dir / "tests" / "weather" / "USA_CO_Denver.intl.AP.725650_TMY3.epw"
    output_dir = measure_dir / "tests" / "workflow_output"
    
    # Create output directory
    output_dir.mkdir(parents=True, exist_ok=True)
    
    # Check if files exist
    if not baseline_model_path.exists():
        print(f"✗ Baseline model not found: {baseline_model_path}")
        sys.exit(1)
    
    if not weather_file.exists():
        print(f"⚠ Weather file not found: {weather_file}")
        print("  EnergyPlus simulations will be skipped")
    
    print(f"Baseline model: {baseline_model_path}")
    print(f"Weather file: {weather_file}")
    print(f"Output directory: {output_dir}\n")
    
    # =========================================================================
    # STEP 1: Process Baseline Model
    # =========================================================================
    print(f"\n{'='*80}")
    print("STEP 1: BASELINE MODEL ANALYSIS")
    print(f"{'='*80}\n")
    
    # Load baseline model
    translator = openstudio.osversion.VersionTranslator()
    baseline_model_opt = translator.loadModel(str(baseline_model_path))
    
    if not baseline_model_opt.is_initialized():
        print("✗ Failed to load baseline model")
        sys.exit(1)
    
    baseline_model = baseline_model_opt.get()
    print(f"✓ Loaded baseline model")
    
    # Extract baseline embodied carbon
    baseline_carbon, baseline_carbon_detail = extract_embodied_carbon(baseline_model)
    print(f"✓ Baseline embodied carbon: {baseline_carbon:.2f} kg CO2eq")
    
    # Run baseline simulation
    baseline_sim_results = None
    if weather_file.exists():
        baseline_sim_dir = output_dir / "baseline_simulation"
        baseline_sim_results = run_energyplus_simulation(
            baseline_model_path, weather_file, baseline_sim_dir
        )
    
    baseline_energy = baseline_sim_results['total_site_energy_GJ'] if baseline_sim_results else 0.0
    print(f"✓ Baseline operational energy: {baseline_energy:.2f} GJ")
    
    # Get baseline cost (could be zero if no RSMeans data)
    baseline_cost = 1000.0  # Placeholder - could use RSMeans
    print(f"✓ Baseline cost: ${baseline_cost:.2f}")
    
    baseline_data = {
        'embodied_carbon': baseline_carbon,
        'operational_energy': baseline_energy,
        'cost': baseline_cost
    }
    
    # =========================================================================
    # STEP 2: Apply Window Enhancement Measure
    # =========================================================================
    print(f"\n{'='*80}")
    print("STEP 2: APPLY WINDOW ENHANCEMENT")
    print(f"{'='*80}\n")
    
    # Create a copy of the model for modification
    modified_model = baseline_model.clone().to_Model()
    
    # Define measure arguments (all 14 required arguments)
    window_args = {
        # Required choice arguments
        'wf_option': 'wood window frame',
        'film_option': 'solar control film',
        'glass_option': 'triple pane clear',
        'gwp_statistic': 'mean',
        
        # Required double arguments
        'caulking_thickness': 0.0127,  # 0.5 inches in meters
        'glass_pane_thickness': 0.003,  # 3mm
        'gap_thickness': 0.013,  # 13mm
        'glass_solar_transmittance': 0.7,
        'glass_front_emissivity': 0.84,
        'glass_back_emissivity': 0.84,
        'length_per_unit': 1.0,  # meters
        
        # Required integer arguments
        'user_num_panes': 3,
        'num_horizontal_dividers': 0,
        'num_vertical_dividers': 0,
        
        # Optional arguments
        'space_infiltration_reduction_percent': 50.0,
        'caulking_option': 'acrylic',
        'weatherstrip_option': 'silicone adhesive smoke gasket',
        'secondary_glazing_option': 'none'
    }
    
    # Apply measure
    success = apply_window_enhancement_measure(modified_model, window_args)
    
    if not success:
        print("✗ Failed to apply window enhancement measure")
        sys.exit(1)
    
    # Save modified model
    modified_model_path = output_dir / "modified_model.osm"
    modified_model.save(str(modified_model_path), True)
    print(f"✓ Saved modified model: {modified_model_path}")
    
    # =========================================================================
    # STEP 3: Process Modified Model
    # =========================================================================
    print(f"\n{'='*80}")
    print("STEP 3: MODIFIED MODEL ANALYSIS")
    print(f"{'='*80}\n")
    
    # Extract modified embodied carbon
    modified_carbon, modified_carbon_detail = extract_embodied_carbon(modified_model)
    print(f"✓ Modified embodied carbon: {modified_carbon:.2f} kg CO2eq")
    
    # Run modified simulation
    modified_sim_results = None
    if weather_file.exists():
        modified_sim_dir = output_dir / "modified_simulation"
        modified_sim_results = run_energyplus_simulation(
            modified_model_path, weather_file, modified_sim_dir
        )
    
    modified_energy = modified_sim_results['total_site_energy_GJ'] if modified_sim_results else 0.0
    print(f"✓ Modified operational energy: {modified_energy:.2f} GJ")
    
    # Get modified cost
    modified_cost = baseline_cost + get_rsmeans_cost("windows")
    print(f"✓ Modified cost: ${modified_cost:.2f}")
    
    modified_data = {
        'embodied_carbon': modified_carbon,
        'operational_energy': modified_energy,
        'cost': modified_cost
    }
    
    # =========================================================================
    # STEP 4: Calculate Deltas
    # =========================================================================
    print(f"\n{'='*80}")
    print("STEP 4: CALCULATE DELTAS")
    print(f"{'='*80}\n")
    
    delta_data = {
        'embodied_carbon': modified_carbon - baseline_carbon,
        'operational_energy': modified_energy - baseline_energy,
        'cost': modified_cost - baseline_cost
    }
    
    print(f"Δ Embodied Carbon: {delta_data['embodied_carbon']:+.2f} kg CO2eq")
    print(f"Δ Operational Energy: {delta_data['operational_energy']:+.2f} GJ")
    print(f"Δ Cost: ${delta_data['cost']:+.2f}")
    
    # =========================================================================
    # STEP 5: Update Optimization Spreadsheet
    # =========================================================================
    updated_excel_path = output_dir / "optimization_updated.xlsx"
    update_optimization_spreadsheet(baseline_data, modified_data, delta_data, updated_excel_path)
    
    # =========================================================================
    # STEP 6: Generate Spider Chart
    # =========================================================================
    spider_chart_path = output_dir / "retrofit_comparison_spider_chart.html"
    generate_spider_chart(baseline_data, modified_data, spider_chart_path)
    
    # =========================================================================
    # FINAL SUMMARY
    # =========================================================================
    print(f"\n\n{'='*80}")
    print("WORKFLOW COMPLETE!")
    print(f"{'='*80}\n")
    print(f"Results saved in: {output_dir}")
    print(f"  - Modified model: modified_model.osm")
    print(f"  - Updated spreadsheet: optimization_updated.xlsx")
    print(f"  - Spider chart: retrofit_comparison_spider_chart.html")
    print(f"\n{'='*80}\n")
    
    return 0


if __name__ == "__main__":
    sys.exit(main())
