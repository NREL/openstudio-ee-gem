import openpyxl, json
wb = openpyxl.load_workbook(r"C:\Users\pshrest2\Downloads\EC3 Query Strings (1).xlsx", data_only=True)
for sn in wb.sheetnames:
    ws = wb[sn]
    print(f"===== SHEET: {sn} (max_row={ws.max_row}, max_col={ws.max_column}) =====")
    for row in ws.iter_rows(values_only=True):
        # skip fully-empty rows
        if all(c is None or (isinstance(c, str) and not c.strip()) for c in row):
            continue
        print(row)
    print()
