"""Build a single-slide PPTX explaining the RSMeans description-matching
scoring scheme as a flow chart.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

OUT = r"c:\All_repos\openstudio-ee-gem\rsmeans_wall_insulation_scoring.pptx"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])

NAVY = RGBColor(0x0B, 0x3D, 0x91)
GREEN = RGBColor(0x1F, 0x7A, 0x3D)
RED = RGBColor(0xB0, 0x2A, 0x2A)
AMBER = RGBColor(0xC8, 0x7A, 0x0E)
GREY = RGBColor(0x55, 0x55, 0x55)
LIGHT = RGBColor(0xF2, 0xF5, 0xFA)
GREEN_FILL = RGBColor(0xE6, 0xF4, 0xEA)
RED_FILL = RGBColor(0xFB, 0xE9, 0xE7)
AMBER_FILL = RGBColor(0xFD, 0xF3, 0xDC)
BLACK = RGBColor(0x10, 0x10, 0x10)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def add_box(left, top, width, height, lines, *, shape=MSO_SHAPE.ROUNDED_RECTANGLE,
            fill=WHITE, border=NAVY, text_color=BLACK, size=11, align=PP_ALIGN.CENTER):
    sh = slide.shapes.add_shape(shape, left, top, width, height)
    sh.fill.solid(); sh.fill.fore_color.rgb = fill
    sh.line.color.rgb = border
    sh.line.width = Pt(1.25)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.06); tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.04); tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = (i == 0 and len(lines) > 1)
        run.font.color.rgb = text_color
    return sh


def add_text(left, top, width, height, lines, *, size=11, bold=False,
             color=BLACK, align=PP_ALIGN.LEFT, fill=None):
    tb = slide.shapes.add_textbox(left, top, width, height)
    if fill is not None:
        tb.fill.solid(); tb.fill.fore_color.rgb = fill
        tb.line.color.rgb = RGBColor(0xCC, 0xD3, 0xE0)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.06); tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.02); tf.margin_bottom = Inches(0.02)
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return tb


def connect(x1, y1, x2, y2, *, color=NAVY, weight=1.5, arrow=True):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    conn.line.color.rgb = color
    conn.line.width = Pt(weight)
    if arrow:
        ln = conn.line._get_or_add_ln()
        for el in ln.findall(qn('a:tailEnd')):
            ln.remove(el)
        tail = ln.makeelement(qn('a:tailEnd'),
                              {'type': 'triangle', 'w': 'med', 'len': 'med'})
        ln.append(tail)
    return conn


# ---- Title ----
add_text(Inches(0.4), Inches(0.18), Inches(12.5), Inches(0.55),
         "RSMeans Description-Match Scoring  -  Flow",
         size=26, bold=True, color=NAVY)
add_text(Inches(0.4), Inches(0.7), Inches(12.5), Inches(0.32),
         "Wall Insulation Retrofit  |  IncreaseInsulationRValueForExteriorWalls  |  resources/call_rsmeans_api.py",
         size=11, color=GREY)

# ---- Top row: linear pipeline ----
row_y = Inches(1.25)
row_h = Inches(0.95)
box_w = Inches(2.25)
gap = Inches(0.18)
start_x = Inches(0.4)

steps = [
    ("Build search term",
     "material name from\ncloned insulation layer"),
    ("Query RSMeans",
     "fetch candidates;\ndrop demolition / disallowed"),
    ("Normalize + tokenize",
     "lowercase, strip punct.;\ntokens > 2 chars"),
    ("Score each candidate",
     "additive rules\n(see panel below)"),
    ("Sort desc.\n& pick best",
     "best score -> decision"),
]

step_centers = []
for i, (head, body) in enumerate(steps):
    x = start_x + i * (box_w + gap)
    add_box(x, row_y, box_w, row_h,
            [head, body],
            fill=LIGHT, border=NAVY, text_color=BLACK, size=10)
    step_centers.append((x + box_w / 2, x, x + box_w))

for i in range(len(steps) - 1):
    cy = row_y + row_h / 2
    _, _, right_i = step_centers[i]
    _, left_next, _ = step_centers[i + 1]
    connect(right_i, cy, left_next, cy, color=NAVY, weight=1.75)

# ---- Decision diamond (centered under last step) ----
last_cx, _, _ = step_centers[-1]
dia_w = Inches(2.6)
dia_h = Inches(1.3)
dia_x = last_cx - dia_w / 2
dia_y = Inches(2.6)
add_box(dia_x, dia_y, dia_w, dia_h,
        ["best score >= 70 ?", "(MIN_ACCEPTABLE_MATCH_SCORE)"],
        shape=MSO_SHAPE.DIAMOND,
        fill=AMBER_FILL, border=AMBER, text_color=BLACK, size=11)

# arrow from last pipeline box bottom to diamond top
connect(last_cx, row_y + row_h, last_cx, dia_y, color=NAVY, weight=1.75)

# ---- Yes outcome (right of diamond) ----
yes_w = Inches(2.55); yes_h = Inches(1.1)
yes_x = dia_x + dia_w + Inches(0.45)
yes_y = dia_y + dia_h / 2 - yes_h / 2
add_box(yes_x, yes_y, yes_w, yes_h,
        ["Use top-scoring line",
         "RSMeans cost line returned\nto measure for cost calc."],
        fill=GREEN_FILL, border=GREEN, text_color=BLACK, size=10)

connect(dia_x + dia_w, dia_y + dia_h / 2, yes_x, yes_y + yes_h / 2,
        color=GREEN, weight=2)
add_text(dia_x + dia_w + Inches(0.02), dia_y + dia_h / 2 - Inches(0.32),
         Inches(0.45), Inches(0.28), "yes", size=11, bold=True, color=GREEN)

# ---- No outcome (left of diamond) ----
no_w = Inches(2.55); no_h = Inches(1.1)
no_x = dia_x - no_w - Inches(0.45)
no_y = dia_y + dia_h / 2 - no_h / 2
add_box(no_x, no_y, no_w, no_h,
        ["Hardcoded fallback ID",
         "Verified 2024-an catalog\nID for retrofit type."],
        fill=RED_FILL, border=RED, text_color=BLACK, size=10)

connect(dia_x, dia_y + dia_h / 2, no_x + no_w, no_y + no_h / 2,
        color=RED, weight=2)
add_text(dia_x - Inches(0.45), dia_y + dia_h / 2 - Inches(0.32),
         Inches(0.4), Inches(0.28), "no", size=11, bold=True, color=RED,
         align=PP_ALIGN.RIGHT)

# ---- Scoring rules panel (bottom) ----
panel_x = Inches(0.4); panel_y = Inches(4.35)
panel_w = Inches(12.5); panel_h = Inches(2.55)
add_box(panel_x, panel_y, panel_w, panel_h, [""],
        shape=MSO_SHAPE.RECTANGLE, fill=WHITE, border=NAVY, size=1)

add_text(panel_x + Inches(0.15), panel_y + Inches(0.05),
         panel_w - Inches(0.3), Inches(0.32),
         "Additive scoring rules applied to each candidate  (final score clamped to 0-100)",
         size=13, bold=True, color=NAVY)

rules_pos = [
    ("+100", "exact normalized match", GREEN),
    ("+60",  "material name is substring of description", GREEN),
    ("+15",  "per overlapping token (>2 chars)", GREEN),
    ("+40",  "'polyisocyanurate' present (polyiso target)", GREEN),
    ("+30",  "'polyiso' present (polyiso target)", GREEN),
    ("+eps", "length tiebreak: min(len,120)/120", GREY),
]
rules_neg = [
    ("-10", "'wall' missing from description (wall target)", RED),
    ("-10", "'roof' missing from description (roof target)", RED),
    ("-20", "'insulation' present but <2 token overlap", RED),
    ("-25", "roof material matched a wall description", RED),
    ("-30", "polyiso target but neither keyword present", RED),
]

col_w = Inches(6.0)
row_h2 = Inches(0.30)
header_y = panel_y + Inches(0.45)

add_text(panel_x + Inches(0.15), header_y, col_w, Inches(0.28),
         "Positive contributions", size=11, bold=True, color=GREEN)
add_text(panel_x + Inches(0.15) + col_w + Inches(0.2), header_y,
         col_w, Inches(0.28),
         "Penalties (cross-type guards)", size=11, bold=True, color=RED)


def draw_rules(items, col_left, top):
    y = top
    for delta, cond, color in items:
        add_text(col_left, y, Inches(0.7), row_h2, delta,
                 size=10.5, bold=True, color=color,
                 align=PP_ALIGN.CENTER, fill=LIGHT)
        add_text(col_left + Inches(0.75), y, col_w - Inches(0.75), row_h2,
                 cond, size=10, color=BLACK)
        y += row_h2


draw_rules(rules_pos, panel_x + Inches(0.15), header_y + Inches(0.32))
draw_rules(rules_neg, panel_x + Inches(0.15) + col_w + Inches(0.2),
           header_y + Inches(0.32))

# Footer
add_text(Inches(0.4), Inches(7.08), Inches(12.5), Inches(0.3),
         "Source: _score_rsmeans_candidate() in call_rsmeans_api.py (L298-351)   |   "
         "Threshold MIN_ACCEPTABLE_MATCH_SCORE at L78",
         size=9, color=GREY)

prs.save(OUT)
print(f"Wrote: {OUT}")
